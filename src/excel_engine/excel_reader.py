"""Low-level Excel file reader supporting .xlsx, .xlsm, and .xls formats."""

from __future__ import annotations

import datetime
from pathlib import Path
from typing import Any

import openpyxl

from src.excel_engine.models import CellValue, RawSheet

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset({".xlsx", ".xlsm", ".xls"})


class UnsupportedFormatError(ValueError):
    """Raised when the file extension is not supported."""


class ExcelReadError(OSError):
    """Raised when an Excel file cannot be opened or parsed."""


class ExcelReader:
    """
    Backend-agnostic Excel reader.

    Uses *openpyxl* for .xlsx/.xlsm and *xlrd* (v2+) for legacy .xls files.
    Returns a list of :class:`RawSheet` objects containing raw cell values
    without any interpretation or transformation.
    """

    def read(self, path: Path) -> list[RawSheet]:
        """
        Read an Excel file and return raw sheet data.

        Args:
            path: Absolute or relative path to the Excel file.

        Returns:
            Ordered list of :class:`RawSheet`, one per worksheet.

        Raises:
            UnsupportedFormatError: File extension is not .xlsx, .xlsm, or .xls.
            ExcelReadError: File is missing or cannot be parsed.
        """
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise UnsupportedFormatError(
                f"Unsupported format {suffix!r}. Supported extensions: {sorted(SUPPORTED_EXTENSIONS)}"
            )
        if not path.exists():
            raise ExcelReadError(f"File not found: {path}")

        if suffix in {".xlsx", ".xlsm"}:
            return self._read_openpyxl(path)
        return self._read_xlrd(path)

    # ------------------------------------------------------------------
    # Private helpers
    # ------------------------------------------------------------------

    def _read_openpyxl(self, path: Path) -> list[RawSheet]:
        """Read an .xlsx / .xlsm file using openpyxl."""
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise ExcelReadError(f"openpyxl could not open {path}: {exc}") from exc

        sheets: list[RawSheet] = []
        try:
            for idx, name in enumerate(wb.sheetnames):
                ws = wb[name]
                state: str = getattr(ws, "sheet_state", "visible")
                is_hidden = state != "visible"
                rows: list[list[CellValue]] = []
                for raw_row in ws.iter_rows(values_only=True):  # type: ignore[union-attr]
                    rows.append([self._coerce_openpyxl(v) for v in raw_row])
                sheets.append(RawSheet(name=name, index=idx, rows=rows, is_hidden=is_hidden))
        finally:
            wb.close()

        return sheets

    def _read_xlrd(self, path: Path) -> list[RawSheet]:
        """Read a legacy .xls file using xlrd."""
        try:
            import xlrd  # type: ignore[import-untyped]
        except ImportError as exc:
            raise ExcelReadError("xlrd is required to read .xls files. Install it with: pip install xlrd") from exc

        try:
            wb = xlrd.open_workbook(str(path))
        except Exception as exc:
            raise ExcelReadError(f"xlrd could not open {path}: {exc}") from exc

        sheets: list[RawSheet] = []
        for idx in range(wb.nsheets):
            ws = wb.sheet_by_index(idx)
            is_hidden = wb.sheet_visibility(idx) != 0
            rows: list[list[CellValue]] = []
            for row_idx in range(ws.nrows):
                row: list[CellValue] = [
                    self._coerce_xlrd(ws.cell(row_idx, col_idx), wb.datemode) for col_idx in range(ws.ncols)
                ]
                rows.append(row)
            sheets.append(RawSheet(name=ws.name, index=idx, rows=rows, is_hidden=is_hidden))

        return sheets

    @staticmethod
    def _coerce_openpyxl(value: Any) -> CellValue:
        """Map an openpyxl raw cell value to a :data:`CellValue`."""
        if value is None:
            return None
        if isinstance(value, bool):
            return value
        if isinstance(value, int):
            return value
        if isinstance(value, float):
            return value
        if isinstance(value, datetime.datetime):
            return value
        if isinstance(value, datetime.date):
            return datetime.datetime(value.year, value.month, value.day)
        return str(value)

    @staticmethod
    def _coerce_xlrd(cell: Any, datemode: int) -> CellValue:
        """Map an xlrd cell to a :data:`CellValue`."""
        try:
            import xlrd  # type: ignore[import-untyped]
        except ImportError:
            return str(cell.value)  # type: ignore[union-attr]

        xtype: int = cell.ctype
        xval: Any = cell.value

        if xtype == xlrd.XL_CELL_EMPTY:
            return None
        if xtype == xlrd.XL_CELL_TEXT:
            return str(xval)
        if xtype == xlrd.XL_CELL_NUMBER:
            if xval == int(xval):
                return int(xval)
            return float(xval)
        if xtype == xlrd.XL_CELL_DATE:
            try:
                return xlrd.xldate_as_datetime(xval, datemode)
            except Exception:
                return str(xval)
        if xtype == xlrd.XL_CELL_BOOLEAN:
            return bool(xval)
        return str(xval)
