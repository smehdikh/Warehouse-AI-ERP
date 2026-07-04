"""Shared fixtures for Excel Engine tests."""

from __future__ import annotations

import datetime
from pathlib import Path

import openpyxl
import pytest

# ---------------------------------------------------------------------------
# Helper
# ---------------------------------------------------------------------------


def make_xlsx(
    tmp_path: Path,
    data: list[list[object]],
    sheet_name: str = "Sheet1",
    filename: str = "test.xlsx",
) -> Path:
    """Write a single-sheet .xlsx file and return its path."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name  # type: ignore[union-attr]
    for row in data:
        ws.append(row)  # type: ignore[union-attr]
    path = tmp_path / filename
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def simple_xlsx(tmp_path: Path) -> Path:
    """Standard tabular file: header + 3 data rows."""
    return make_xlsx(
        tmp_path,
        [
            ["ID", "Name", "Price", "Quantity"],
            [1, "Widget A", 9.99, 100],
            [2, "Widget B", 19.99, 50],
            [3, "Widget C", 4.99, 200],
        ],
    )


@pytest.fixture
def empty_xlsx(tmp_path: Path) -> Path:
    """Workbook with no data rows."""
    return make_xlsx(tmp_path, [], filename="empty.xlsx")


@pytest.fixture
def blanks_before_header_xlsx(tmp_path: Path) -> Path:
    """Two blank rows, then the header, then data."""
    return make_xlsx(
        tmp_path,
        [
            [None, None, None],
            [None, None, None],
            ["SKU", "Description", "Unit Price"],
            ["A001", "Bolt M8", 0.15],
            ["A002", "Nut M8", 0.08],
        ],
        filename="blanks_before_header.xlsx",
    )


@pytest.fixture
def multi_sheet_xlsx(tmp_path: Path) -> Path:
    """Workbook with two data sheets and one empty sheet."""
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Products"  # type: ignore[union-attr]
    for row in [["ID", "Name"], [1, "Alpha"], [2, "Beta"]]:
        ws1.append(row)  # type: ignore[union-attr]

    ws2 = wb.create_sheet("Suppliers")
    for row in [["Code", "Company"], ["S1", "Acme"], ["S2", "Beta Corp"]]:
        ws2.append(row)

    wb.create_sheet("EmptySheet")  # intentionally empty

    path = tmp_path / "multi_sheet.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def duplicates_xlsx(tmp_path: Path) -> Path:
    """File with duplicate rows."""
    return make_xlsx(
        tmp_path,
        [
            ["ID", "Name"],
            [1, "Alpha"],
            [1, "Alpha"],  # duplicate
            [2, "Beta"],
        ],
        filename="duplicates.xlsx",
    )


@pytest.fixture
def mixed_types_xlsx(tmp_path: Path) -> Path:
    """Columns with mixed types."""
    return make_xlsx(
        tmp_path,
        [
            ["ref", "value", "date_col"],
            ["A", 1, datetime.datetime(2024, 1, 1)],
            ["B", 2.5, datetime.datetime(2024, 1, 2)],
            ["C", 3, datetime.datetime(2024, 1, 3)],
        ],
        filename="mixed_types.xlsx",
    )
